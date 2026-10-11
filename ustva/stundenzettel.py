"""Stundenzettel: feste Stunden je Monat an einem festen Wochentag, als Excel-Datei mit einem Blatt je Monat.

Das Layout folgt der Datei, die bisher von Hand gepflegt wurde (Name und „Stundenzettel“ im Kopf,
Monat, Zeilen mit Wochentag, Datum und Stunden, Summe, Feld für Datum und Unterschrift).
Vorgeschlagen werden die Arbeitstage des Monats (Standard: alle Mittwoche ohne Feiertag), auf die
sich die Monatsstunden in halben Stunden verteilen. Vor dem Eintragen lassen sich die Tage anpassen,
etwa in der Urlaubszeit.
"""

from __future__ import annotations

import calendar
import io
import json
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation

from .db import Datenbank
from .fristen import MONATSNAMEN, WOCHENTAGE, feiertage

EINSTELLUNG = "stundenzettel"
SCHRIFT = "Aptos Narrow"
MIN_ZEILEN = 7  # so viele Zeilen hat die Vorlage, auch wenn weniger Tage eingetragen sind


@dataclass
class Einstellungen:
    name: str = ""  # steht oben links, z. B. „Nachname Vorname“
    stunden: Decimal = Decimal(20)  # je Monat
    wochentag: int = 2  # 0 = Montag, 2 = Mittwoch
    ort: str = ""  # steht mit dem Datum über der Unterschriftslinie

    def to_dict(self) -> dict:
        return {"name": self.name, "stunden": str(self.stunden), "wochentag": self.wochentag, "ort": self.ort}

    @classmethod
    def aus_dict(cls, d: dict) -> Einstellungen:
        e = cls()
        e.name = str(d.get("name") or "").strip()
        e.ort = str(d.get("ort") or "").strip()
        if d.get("stunden") not in (None, ""):
            try:
                e.stunden = Decimal(str(d["stunden"]).replace(",", "."))
            except InvalidOperation as ex:
                raise ValueError(f"Keine Zahl: {d['stunden']}") from ex
        if d.get("wochentag") not in (None, ""):
            e.wochentag = int(d["wochentag"])
        return e


def laden(db: Datenbank) -> Einstellungen:
    roh = db.einstellungen().get(EINSTELLUNG)
    return Einstellungen.aus_dict(json.loads(roh)) if roh else Einstellungen()


def speichern(db: Datenbank, e: Einstellungen) -> Einstellungen:
    if not Decimal(0) < e.stunden <= 300:
        raise ValueError("Die Stunden je Monat müssen zwischen 0 und 300 liegen.")
    if not 0 <= e.wochentag <= 6:
        raise ValueError("Wochentag ungültig.")
    db.einstellung_setzen(EINSTELLUNG, json.dumps(e.to_dict(), ensure_ascii=False))
    return e


def verteilen(stunden: Decimal, tage: list[date]) -> list[tuple[date, Decimal]]:
    """Verteilt die Stunden in halben Stunden möglichst gleichmäßig, die ersten Tage bekommen den Rest."""
    if not tage:
        return []
    einheiten = int(stunden * 2)
    je_tag, rest = divmod(einheiten, len(tage))
    return [(t, Decimal(je_tag + (1 if i < rest else 0)) / 2) for i, t in enumerate(sorted(tage))]


def vorschlag(e: Einstellungen, jahr: int, monat: int, bundesland: str) -> dict:
    """Alle Arbeitstage (fester Wochentag) des Monats ohne Feiertage, mit den Monatsstunden verteilt."""
    frei = feiertage(jahr, bundesland)
    tag = date(jahr, monat, 1)
    tag += timedelta(days=(e.wochentag - tag.weekday()) % 7)
    tage, uebersprungen = [], []
    while tag.month == monat:
        if tag in frei:
            uebersprungen.append({"datum": tag.isoformat(), "name": frei[tag]})
        else:
            tage.append(tag)
        tag += timedelta(days=7)
    return {
        "tage": [{"datum": d.isoformat(), "stunden": str(h)} for d, h in verteilen(e.stunden, tage)],
        "feiertage": uebersprungen,
    }


def letzter_monat(jahr: int, heute: date, belegt: set[int]) -> int:
    """Bis zu welchem Monat die Datei Blätter hat: nichts in der Zukunft, der laufende Monat nur mit Einträgen."""
    if jahr < heute.year:
        return 12
    if jahr > heute.year:
        return 0
    return heute.month if heute.month in belegt else heute.month - 1


def monate(db: Datenbank, jahr: int, heute: date) -> list[dict]:
    nach_monat: dict[int, list[tuple[date, Decimal]]] = {}
    for d, h in db.stundenzettel(jahr):
        nach_monat.setdefault(d.month, []).append((d, h))
    bis = max(letzter_monat(jahr, heute, set(nach_monat)), 1 if jahr <= heute.year else 0)
    return [{
        "zeitraum": f"{jahr:04d}-{m:02d}",
        "name": MONATSNAMEN[m - 1],
        "tage": [{"datum": d.isoformat(), "stunden": str(h)} for d, h in nach_monat.get(m, [])],
        "stunden": str(sum((h for _, h in nach_monat.get(m, [])), Decimal(0))),
    } for m in range(1, bis + 1)]


def monat_pruefen(zeitraum: str, tage: list[tuple[date, Decimal]], heute: date) -> None:
    jahr, monat = int(zeitraum[:4]), int(zeitraum[5:7])
    if (jahr, monat) > (heute.year, heute.month):
        raise ValueError("Monate in der Zukunft lassen sich noch nicht eintragen.")
    gesehen: set[date] = set()
    for d, h in tage:
        if (d.year, d.month) != (jahr, monat):
            raise ValueError(f"Der {d:%d.%m.%Y} liegt nicht im {MONATSNAMEN[monat - 1]} {jahr}.")
        if d in gesehen:
            raise ValueError(f"Der {d:%d.%m.%Y} steht doppelt drin.")
        if not Decimal(0) < h <= 24:
            raise ValueError("Stunden je Tag zwischen 0 und 24 angeben.")
        gesehen.add(d)


# Excel-Ausgabe

def _zahl(h: Decimal) -> int | float:
    return int(h) if h == h.to_integral_value() else float(h)


def unterschriftsdatum(jahr: int, monat: int) -> date:
    """Unterschrieben wird zum Monatsende."""
    return date(jahr, monat, calendar.monthrange(jahr, monat)[1])


def _blatt(wb, e: Einstellungen, jahr: int, monat: int, tage: list[tuple[date, Decimal]]) -> None:
    ws = wb.add_worksheet(MONATSNAMEN[monat - 1])

    def fmt(**kw):
        return wb.add_format({"font_name": SCHRIFT, "font_size": 14, **kw})

    kopf_links = fmt(bold=True, font_size=16, align="left")
    kopf_rechts = fmt(bold=True, font_size=16, align="right")
    normal = fmt()
    fett_rechts = fmt(bold=True, align="right")
    rechts = fmt(align="right")
    datum = fmt(num_format="dd/mm/yyyy")
    letzte_stunden = fmt(bottom=6)  # doppelte Linie unter der letzten Zeile
    summe = fmt(bottom=1)
    unterschrift_links = fmt(font_size=9, align="left", top=1)
    unterschrift_mitte = fmt(font_size=9, top=1)
    unterschrift_rechts = fmt(font_size=9, align="right", top=1)

    ws.set_column(0, 1, 13.5)
    ws.set_column(2, 2, 16.5)
    ws.set_column(3, 4, 13.5)
    ws.set_row(0, 22)
    ws.merge_range(0, 0, 0, 1, e.name, kopf_links)
    ws.merge_range(0, 2, 0, 4, "Stundenzettel", kopf_rechts)
    ws.set_row(2, 20)
    ws.write(2, 0, "Monat:", normal)
    ws.write(2, 2, MONATSNAMEN[monat - 1], fett_rechts)
    ws.write(2, 4, "Stunden", rechts)

    erste = 4
    zeilen = max(MIN_ZEILEN, len(tage))
    for i in range(zeilen):
        r = erste + i
        ws.set_row(r, 20)
        letzte = i == zeilen - 1
        if i < len(tage):
            d, h = tage[i]
            ws.write(r, 0, WOCHENTAGE[d.weekday()], normal)
            ws.write_datetime(r, 2, datetime(d.year, d.month, d.day), datum)
            ws.write_number(r, 4, _zahl(h), letzte_stunden if letzte else normal)
        elif letzte:
            ws.write_blank(r, 4, None, letzte_stunden)

    g = erste + zeilen + 1
    ws.set_row(g, 20)
    ws.write(g, 0, "Gesamt:", summe)
    for c in (1, 2, 3):
        ws.write_blank(g, c, None, summe)
    gesamt = sum((h for _, h in tage), Decimal(0))
    ws.write_formula(g, 4, f"=SUM(E{erste + 1}:E{erste + zeilen})", summe, _zahl(gesamt))

    u = g + 4
    if tage:
        ort_datum = f"{unterschriftsdatum(jahr, monat):%d.%m.%Y}"
        ws.write(u - 1, 2, f"{e.ort}, {ort_datum}" if e.ort else ort_datum, fmt(font_size=12))
    ws.write(u, 2, "Datum", unterschrift_links)
    ws.write_blank(u, 3, None, unterschrift_mitte)
    ws.write(u, 4, "Unterschrift", unterschrift_rechts)

    ws.set_paper(9)
    ws.set_portrait()
    ws.set_margins(left=0.7, right=0.7, top=0.79, bottom=0.79)
    ws.print_area(0, 0, u, 4)


def xlsx(e: Einstellungen, jahr: int, eintraege: list[tuple[date, Decimal]], bis: int) -> bytes:
    """Ein Blatt je Monat von Januar bis ``bis``."""
    import xlsxwriter

    puffer = io.BytesIO()
    wb = xlsxwriter.Workbook(puffer, {"in_memory": True})
    for m in range(1, bis + 1):
        _blatt(wb, e, jahr, m, [(d, h) for d, h in eintraege if d.month == m])
    if bis:
        wb.worksheets()[bis - 1].activate()
    wb.close()
    return puffer.getvalue()


def dateiname(e: Einstellungen, jahr: int) -> str:
    return " ".join(t for t in (e.name, "Stundenzettel", str(jahr)) if t) + ".xlsx"


# Übernahme aus der bisherigen Excel-Datei

@dataclass
class ImportErgebnis:
    tage: int = 0
    monate: list[str] = field(default_factory=list)
    uebersprungen: list[str] = field(default_factory=list)  # schon erfasst, bleibt unverändert
    andere_jahre: list[str] = field(default_factory=list)  # Blätter mit Daten aus einem anderen Jahr
    name: str = ""

    def to_dict(self) -> dict:
        return {"tage": self.tage, "monate": self.monate, "uebersprungen": self.uebersprungen,
                "andere_jahre": self.andere_jahre, "name": self.name}


def importieren(db: Datenbank, inhalt: bytes, jahr: int) -> ImportErgebnis:
    """Übernimmt die Tage eines Jahres; Zeilen mit einem Datum aus einem anderen Jahr bleiben draußen."""
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(inhalt), data_only=True)
    except Exception as ex:  # openpyxl wirft je nach Defekt sehr unterschiedliche Fehler
        raise ValueError("Die Datei lässt sich nicht als Excel-Datei (.xlsx) lesen.") from ex

    ergebnis = ImportErgebnis()
    nach_monat: dict[str, list[tuple[date, Decimal]]] = {}
    erkannt = False
    for ws in wb.worksheets:
        if str(ws["C1"].value or "").strip().lower() != "stundenzettel":
            continue
        erkannt = True
        ergebnis.name = ergebnis.name or str(ws["A1"].value or "").strip()
        fremd: set[int] = set()
        for zeile in ws.iter_rows(min_row=4, max_col=5, values_only=True):
            if str(zeile[0] or "").strip().lower().startswith("gesamt"):
                break
            d, h = zeile[2], zeile[4]
            if not isinstance(d, datetime | date) or not isinstance(h, int | float) or h <= 0:
                continue
            d = d.date() if isinstance(d, datetime) else d
            if d.year != jahr:
                fremd.add(d.year)
                continue
            nach_monat.setdefault(f"{d:%Y-%m}", []).append((d, Decimal(str(h))))
        if fremd:
            ergebnis.andere_jahre.append(f"{ws.title} ({', '.join(map(str, sorted(fremd)))})")
    if not erkannt:
        raise ValueError("In der Datei steht kein Stundenzettel im erwarteten Format.")

    vorhanden = {f"{d:%Y-%m}" for d, _ in db.stundenzettel(jahr)}
    for zr, tage in sorted(nach_monat.items()):
        if zr in vorhanden:
            ergebnis.uebersprungen.append(zr)
            continue
        db.stundenzettel_monat_setzen(zr, sorted(dict(tage).items()))
        ergebnis.monate.append(zr)
        ergebnis.tage += len(dict(tage))

    e = laden(db)
    if ergebnis.name and not e.name:
        e.name = ergebnis.name
        speichern(db, e)
    return ergebnis
