"""Tätigkeitsnachweis: Stunden je Tag erfassen und monatlich als Excel-Datei ausgeben.

Das Layout folgt der Excel-Vorlage, die Marcel bisher von Hand pflegt (ein Blatt
je Monat, Spalten Datum bis Gesamt, Summe, geleistete Stunden, Unterschriftsfeld).
Was sich nicht täglich ändert (Kopf, Name, Stundensatz, Aufträge, Auswahllisten),
steht in der Vorlage. Sie liegt als JSON in der Einstellungen-Tabelle und lässt
sich aus einer bestehenden Excel-Datei übernehmen.
"""

from __future__ import annotations

import io
import json
import math
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation

from .db import Datenbank
from .fristen import feiertage
from .models import Taetigkeit

VORLAGE_SCHLUESSEL = "taetigkeit_vorlage"
MONATE = ["Januar", "Februar", "März", "April", "Mai", "Juni", "Juli", "August", "September", "Oktober",
          "November", "Dezember"]
SPALTEN = ["Datum", "Inventarnummer", "Inventar", "Auftragsnummer", "Auftrag", "Ressource", "Rolle", "Ort",
           "Referenz", "Leistungsart", "Tätigkeit", "Satz", "Menge", "Gesamt"]
BREITEN = [11.5, 16.5, 24.5, 28.16, 23.16, 18.16, 19.5, 10.5, 11.16, 15.83, 42.5, 11.5, 11.5, 16.5]


@dataclass
class Auftrag:
    name: str
    nummer: str = ""


@dataclass
class Vorlage:
    titel: str = "Tätigkeitsnachweis"
    name: str = ""
    vorname: str = ""
    kopf: list[str] = field(default_factory=list)  # rechts oben, z. B. Firma und Anschrift
    stundensatz: Decimal | None = None
    stunden: Decimal = Decimal(8)  # Vorschlag für einen neuen Tag
    auftraege: list[Auftrag] = field(default_factory=list)
    rollen: list[str] = field(default_factory=lambda: ["Entwickler"])
    orte: list[str] = field(default_factory=lambda: ["Remote", "vor Ort"])
    leistungsarten: list[str] = field(default_factory=lambda: ["Entwicklung"])

    @property
    def ressource(self) -> str:
        return ", ".join(t for t in (self.name, self.vorname) if t)

    def auftragsnummer(self, auftrag: str) -> str:
        return next((a.nummer for a in self.auftraege if a.name == auftrag), "")

    def to_dict(self) -> dict:
        d = asdict(self)
        d["stundensatz"] = None if self.stundensatz is None else str(self.stundensatz)
        d["stunden"] = str(self.stunden)
        return d

    @classmethod
    def aus_dict(cls, d: dict) -> Vorlage:
        def texte(werte) -> list[str]:
            return [str(w).strip() for w in werte or [] if str(w).strip()]

        v = cls()
        v.titel = str(d.get("titel") or v.titel).strip()
        v.name = str(d.get("name") or "").strip()
        v.vorname = str(d.get("vorname") or "").strip()
        v.kopf = texte(d.get("kopf"))
        v.stundensatz = _dezimal(d.get("stundensatz"))
        v.stunden = _dezimal(d.get("stunden")) or v.stunden
        v.auftraege = [Auftrag(str(a.get("name", "")).strip(), str(a.get("nummer") or "").strip())
                       for a in d.get("auftraege") or [] if str(a.get("name", "")).strip()]
        for liste in ("rollen", "orte", "leistungsarten"):
            if liste in d:
                setattr(v, liste, texte(d[liste]))
        return v


def _dezimal(wert) -> Decimal | None:
    if wert in (None, ""):
        return None
    try:
        return Decimal(str(wert).replace(",", "."))
    except InvalidOperation as e:
        raise ValueError(f"Keine Zahl: {wert}") from e


def vorlage_laden(db: Datenbank) -> Vorlage:
    roh = db.einstellungen().get(VORLAGE_SCHLUESSEL)
    return Vorlage.aus_dict(json.loads(roh)) if roh else Vorlage()


def vorlage_speichern(db: Datenbank, v: Vorlage) -> Vorlage:
    if v.stundensatz is not None and v.stundensatz < 0:
        raise ValueError("Der Stundensatz darf nicht negativ sein.")
    if not Decimal(0) < v.stunden <= 24:
        raise ValueError("Die Standard-Stunden müssen zwischen 0 und 24 liegen.")
    db.einstellung_setzen(VORLAGE_SCHLUESSEL, json.dumps(v.to_dict(), ensure_ascii=False))
    return v


def satz(t: Taetigkeit, v: Vorlage) -> Decimal | None:
    return t.satz if t.satz is not None else v.stundensatz


def betrag(t: Taetigkeit, v: Vorlage) -> Decimal | None:
    """Ungerundet wie die Formel in der Excel-Datei; gerundet wird erst die Anzeige."""
    s = satz(t, v)
    return None if s is None else s * t.stunden


# Vorschläge für die Erfassung

def _monatstage(jahr: int, monat: int) -> list[date]:
    tag, tage = date(jahr, monat, 1), []
    while tag.month == monat:
        tage.append(tag)
        tag += timedelta(days=1)
    return tage


def _werktag(tag: date, bundesland: str) -> bool:
    return tag.weekday() < 5 and tag not in feiertage(tag.year, bundesland)


def fehlende_tage(eintraege: list[Taetigkeit], jahr: int, monat: int, heute: date, bundesland: str) -> list[date]:
    """Werktage des Monats bis heute ohne Eintrag (Urlaub erkennt die App natürlich nicht)."""
    belegt = {t.datum for t in eintraege}
    return [t for t in _monatstage(jahr, monat) if t <= heute and _werktag(t, bundesland) and t not in belegt]


def vorschlag(db: Datenbank, v: Vorlage, jahr: int, monat: int, bundesland: str) -> dict:
    """Vorbelegung für den nächsten Eintrag: Felder vom letzten Eintrag, Datum der nächste freie Werktag."""
    im_monat = db.taetigkeiten(f"{jahr:04d}-{monat:02d}")
    alle = im_monat or db.taetigkeiten()
    letzter = alle[-1] if alle else None
    belegt = {t.datum for t in im_monat}
    ab = im_monat[-1].datum if im_monat else date(jahr, monat, 1) - timedelta(days=1)
    frei = [t for t in _monatstage(jahr, monat) if t > ab and t not in belegt and _werktag(t, bundesland)]
    datum = frei[0] if frei else (im_monat[-1].datum if im_monat else date(jahr, monat, 1))
    return {
        "datum": datum.isoformat(),
        "auftrag": letzter.auftrag if letzter else (v.auftraege[0].name if v.auftraege else ""),
        "referenz": letzter.referenz if letzter else "",
        "taetigkeit": letzter.taetigkeit if letzter else "",
        "stunden": str(v.stunden),
        "rolle": letzter.rolle if letzter else (v.rollen[0] if v.rollen else ""),
        "ort": letzter.ort if letzter else (v.orte[0] if v.orte else ""),
        "leistungsart": letzter.leistungsart if letzter else (v.leistungsarten[0] if v.leistungsarten else ""),
    }


def fruehere_texte(db: Datenbank, grenze: int = 300) -> list[dict]:
    """Bisherige Tätigkeiten (neueste zuerst, ohne Doppelte) zum Wiederverwenden."""
    gesehen: set[str] = set()
    texte = []
    for t in reversed(db.taetigkeiten()):
        schluessel = t.taetigkeit.strip()
        if not schluessel or schluessel in gesehen:
            continue
        gesehen.add(schluessel)
        texte.append({"taetigkeit": schluessel, "auftrag": t.auftrag, "referenz": t.referenz})
        if len(texte) >= grenze:
            break
    return texte


# Excel-Ausgabe

def _referenz(wert: str):
    """Ticketnummern stehen in der Vorlage als Zahl."""
    return int(wert) if wert.isdigit() and len(wert) < 15 else wert


def _zeilenhoehe(texte: list[tuple[str, float]]) -> float:
    """Grobe Höhe für umbrochenen Text; Excel passt Zeilen beim Öffnen nicht selbst an."""
    zeilen = 1
    for text, breite in texte:
        zeichen = max(int(breite * 0.95), 1)
        zeilen = max(zeilen, sum(max(1, math.ceil(len(teil) / zeichen)) for teil in text.strip().split("\n")))
    return max(30.0, 15.0 * zeilen)


def nachweis_xlsx(v: Vorlage, eintraege: list[Taetigkeit], jahr: int, monat: int) -> bytes:
    """Ein Monatsblatt im Layout des bisherigen Tätigkeitsnachweises."""
    import xlsxwriter

    puffer = io.BytesIO()
    wb = xlsxwriter.Workbook(puffer, {"in_memory": True})
    ws = wb.add_worksheet(MONATE[monat - 1])

    basis = {"font_name": "Arial", "font_size": 11, "align": "left"}
    blau = {**basis, "font_color": "#0000FF", "valign": "top", "border": 1}

    def fmt(**kw):
        return wb.add_format({**basis, **kw})

    titel = fmt(bold=True, font_size=14)
    normal = fmt()
    rechts = fmt(font_size=9, align="right")
    kopf = fmt(bg_color="#B7DEE8", border=1, top=2)
    kopf_links = fmt(bg_color="#B7DEE8", border=1, top=2, left=2)
    kopf_rechts = fmt(bg_color="#B7DEE8", border=1, top=2, right=2)
    zelle = wb.add_format(blau)
    zelle_links = wb.add_format({**blau, "left": 2, "num_format": "dd/mm/yy;@"})
    zelle_text = wb.add_format({**blau, "text_wrap": True})
    zelle_zahl = wb.add_format({**blau, "num_format": "0"})
    zelle_satz = wb.add_format({**blau, "num_format": '#,##0.00" €"'})
    zelle_menge = wb.add_format({**blau, "num_format": "0.00"})
    zelle_gesamt = wb.add_format({**blau, "num_format": '#,##0.00" €"', "right": 2})
    summe = fmt(top=2, bottom=2)
    summe_links = fmt(top=2, bottom=2, left=2)
    summe_menge = fmt(top=2, bottom=2, num_format="0.00")
    summe_betrag = fmt(top=2, bottom=2, right=2, num_format='#,##0.00" €"')
    stunden_titel = fmt(bold=True, border=2, valign="top")
    stunden_wert = fmt(bold=True, border=2, align="center", valign="vcenter", bg_color="#B7DEE8",
                       num_format="0.00", font_size=14)
    unterschrift = fmt(bold=True)

    for i, breite in enumerate(BREITEN):
        ws.set_column(i, i, breite)

    # Kopf (Zeilen wie in der Vorlage, xlsxwriter zählt ab 0)
    ws.set_row(0, 18)
    ws.write(0, 0, v.titel, titel)
    for i, zeile in enumerate(v.kopf[:5]):
        ws.write(1 + i, 13, zeile, rechts)
    ws.write(2, 0, "Monat:", normal)
    ws.write(2, 1, MONATE[monat - 1], normal)
    ws.write(3, 0, "Name:", normal)
    ws.write(3, 1, v.name, normal)
    ws.write(4, 0, "Vorname:", normal)
    ws.write(4, 1, v.vorname, normal)

    erste = 7
    for i, spalte in enumerate(SPALTEN):
        ws.write(6, i, spalte, kopf_links if i == 0 else kopf_rechts if i == len(SPALTEN) - 1 else kopf)

    eintraege = sorted(eintraege, key=lambda t: (t.datum, t.id or 0))
    gesamt_stunden = Decimal(0)
    gesamt_betrag = Decimal(0)
    for n, t in enumerate(eintraege):
        r = erste + n
        excel = r + 1
        s = satz(t, v)
        b = betrag(t, v) or Decimal(0)
        gesamt_stunden += t.stunden
        gesamt_betrag += b
        ws.set_row(r, _zeilenhoehe([(t.taetigkeit, BREITEN[10]), (t.auftrag, BREITEN[4])]))
        ws.write_datetime(r, 0, datetime(t.datum.year, t.datum.month, t.datum.day), zelle_links)
        ws.write_blank(r, 1, None, zelle)
        ws.write_blank(r, 2, None, zelle)
        ws.write(r, 3, v.auftragsnummer(t.auftrag), zelle)
        ws.write(r, 4, t.auftrag, zelle_text)
        ws.write(r, 5, v.ressource, zelle)
        ws.write(r, 6, t.rolle, zelle)
        ws.write(r, 7, t.ort, zelle)
        ws.write(r, 8, _referenz(t.referenz), zelle_zahl)
        ws.write(r, 9, t.leistungsart, zelle)
        ws.write(r, 10, t.taetigkeit.strip(), zelle_text)
        if s is None:
            ws.write_blank(r, 11, None, zelle_satz)
        else:
            ws.write_number(r, 11, float(s), zelle_satz)
        ws.write_number(r, 12, float(t.stunden), zelle_menge)
        ws.write_formula(r, 13, f"=L{excel}*M{excel}", zelle_gesamt, float(b))

    letzte = erste + max(len(eintraege), 1) - 1  # mindestens eine (leere) Zeile, damit die Summen gültig sind
    if not eintraege:
        for i in range(len(SPALTEN)):
            ws.write_blank(erste, i, None, zelle_links if i == 0 else zelle)
    s = letzte + 1  # Summenzeile
    bereich = f"{erste + 1}:{letzte + 1}"
    oben, unten = bereich.split(":")
    ws.write(s, 0, "Gesamt", summe_links)
    for i in range(1, 12):
        ws.write_blank(s, i, None, summe)
    ws.write_formula(s, 12, f"=SUM(M{oben}:M{unten})", summe_menge, float(gesamt_stunden))
    ws.write_formula(s, 13, f"=SUM(N{oben}:N{unten})", summe_betrag, float(gesamt_betrag))

    g = s + 2  # geleistete Stunden
    ws.merge_range(g, 0, g + 1, 13, "geleistete Stunden:", stunden_titel)
    ws.merge_range(g + 2, 0, g + 5, 13, "", stunden_wert)
    ws.write_formula(g + 2, 0, f"=M{s + 1}", stunden_wert, float(gesamt_stunden))

    u = g + 9  # Unterschrift Kunde
    ws.write(u, 10, "Ort:", unterschrift)
    ws.write(u, 13, "Datum:", unterschrift)
    ws.write(u + 2, 10, "Hiermit bestätige ich die ordnungsgemäße", normal)
    ws.write(u + 3, 10, "Leistungserbringung", normal)
    ws.write(u + 5, 10, "Unterschrift Kunde:", unterschrift)

    ws.autofilter(6, 0, letzte, 13)
    ws.freeze_panes(7, 0)
    ws.set_paper(9)
    ws.set_landscape()
    ws.fit_to_pages(1, 0)
    ws.print_area(0, 0, u + 5, 13)
    ws.set_margins(left=0.6, right=0.6, top=0.6, bottom=0.8)
    ws.set_footer("&R&8Seite &P von &N")
    wb.close()
    return puffer.getvalue()


def dateiname(jahr: int, monat: int) -> str:
    return f"Tätigkeitsnachweis {MONATE[monat - 1]} {jahr}.xlsx"


# Übernahme aus einer bestehenden Excel-Datei

@dataclass
class ImportErgebnis:
    eintraege: int = 0
    monate: list[str] = field(default_factory=list)
    uebersprungen: list[str] = field(default_factory=list)  # Monate, die in der App schon Einträge haben
    vorlage: list[str] = field(default_factory=list)  # übernommene Teile der Vorlage

    def to_dict(self) -> dict:
        return asdict(self)


def _text(wert) -> str:
    if wert is None:
        return ""
    if isinstance(wert, float) and wert.is_integer():
        wert = int(wert)
    return str(wert).strip()


def _tabellen(ws) -> list[list[list]]:
    """Inhalt aller als Tabelle formatierten Bereiche eines Blatts (Kopfzeile zuerst)."""
    ergebnis = []
    for tabelle in ws.tables.values():
        ergebnis.append([[c.value for c in zeile] for zeile in ws[tabelle.ref]])
    return ergebnis


def _werte_lesen(ws, v: Vorlage, ergebnis: ImportErgebnis) -> None:
    for zeilen in _tabellen(ws):
        kopf = [_text(k).lower() for k in zeilen[0]]
        daten = zeilen[1:]
        if "auftrag" in kopf and "auftragsnummer" in kopf:
            a, n = kopf.index("auftrag"), kopf.index("auftragsnummer")
            auftraege = [Auftrag(_text(z[a]), _text(z[n])) for z in daten if _text(z[a])]
            if auftraege:
                v.auftraege = auftraege
                ergebnis.vorlage.append(f"{len(auftraege)} Aufträge")
        for name, liste in (("rolle", "rollen"), ("ort", "orte"), ("leistungsart", "leistungsarten")):
            if name in kopf:
                werte = [_text(z[kopf.index(name)]) for z in daten if _text(z[kopf.index(name)])]
                if werte:
                    setattr(v, liste, werte)
                    ergebnis.vorlage.append(f"{len(werte)} {liste.capitalize()}")


def _kopf_lesen(ws, v: Vorlage, ergebnis: ImportErgebnis) -> None:
    if titel := _text(ws["A1"].value):
        v.titel = titel
    kopf = [_text(ws.cell(row=r, column=14).value) for r in range(1, 7)]
    if any(kopf):
        v.kopf = [k for k in kopf if k]
    for zeile in range(2, 7):
        feld = _text(ws.cell(row=zeile, column=1).value).lower().rstrip(":")
        wert = _text(ws.cell(row=zeile, column=2).value)
        if feld == "name" and wert:
            v.name = wert
        elif feld == "vorname" and wert:
            v.vorname = wert
    if "Kopf" not in ergebnis.vorlage:
        ergebnis.vorlage.append("Kopf")


def _eintraege_lesen(ws) -> tuple[list[Taetigkeit], Decimal | None]:
    kopfzeile, spalten = None, {}
    for zeile in ws.iter_rows(min_row=1, max_row=20):
        namen = [_text(c.value) for c in zeile]
        if "Datum" in namen and "Tätigkeit" in namen:
            kopfzeile = zeile[0].row
            spalten = {n: i for i, n in enumerate(namen) if n}
            break
    if kopfzeile is None:
        return [], None

    def wert(z, name):
        i = spalten.get(name)
        return z[i] if i is not None and i < len(z) else None

    eintraege, letzter_satz = [], None
    for z in ws.iter_rows(min_row=kopfzeile + 1, values_only=True):
        datum = wert(z, "Datum")
        if _text(datum).lower() == "gesamt":
            break
        if not isinstance(datum, datetime | date):
            continue
        menge = wert(z, "Menge")
        if not isinstance(menge, int | float) or menge <= 0:
            continue
        satz_ = wert(z, "Satz")
        satz_ = Decimal(str(satz_)) if isinstance(satz_, int | float) else None
        letzter_satz = satz_ or letzter_satz
        eintraege.append(Taetigkeit(
            datum=datum.date() if isinstance(datum, datetime) else datum,
            stunden=Decimal(str(menge)),
            taetigkeit=_text(wert(z, "Tätigkeit")),
            auftrag=_text(wert(z, "Auftrag")),
            referenz=_text(wert(z, "Referenz")),
            satz=satz_,
            rolle=_text(wert(z, "Rolle")),
            ort=_text(wert(z, "Ort")),
            leistungsart=_text(wert(z, "Leistungsart")),
        ))
    return eintraege, letzter_satz


def importieren(db: Datenbank, inhalt: bytes) -> ImportErgebnis:
    """Übernimmt Vorlage und Einträge aus einer Excel-Datei im bisherigen Format.

    Monate, für die in der App schon Einträge stehen, bleiben unverändert.
    """
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(inhalt), data_only=True)
    except Exception as e:  # openpyxl wirft je nach Defekt sehr unterschiedliche Fehler
        raise ValueError("Die Datei lässt sich nicht als Excel-Datei (.xlsx) lesen.") from e

    v = vorlage_laden(db)
    ergebnis = ImportErgebnis()
    nach_monat: dict[str, list[Taetigkeit]] = {}
    satz_gefunden = None
    for ws in wb.worksheets:
        if any(_text(k).lower() == "auftragsnummer" for z in _tabellen(ws) for k in z[0]):
            _werte_lesen(ws, v, ergebnis)
            continue
        eintraege, satz_ = _eintraege_lesen(ws)
        if not eintraege and not _text(ws["A1"].value).lower().startswith("tätigkeitsnachweis"):
            continue
        _kopf_lesen(ws, v, ergebnis)
        satz_gefunden = satz_ or satz_gefunden
        for t in eintraege:
            nach_monat.setdefault(f"{t.datum:%Y-%m}", []).append(t)

    if satz_gefunden is not None:
        v.stundensatz = satz_gefunden
        ergebnis.vorlage.append("Stundensatz")
    if not ergebnis.vorlage and not nach_monat:
        raise ValueError("In der Datei steht kein Tätigkeitsnachweis im erwarteten Format.")

    vorhanden = {f"{t.datum:%Y-%m}" for t in db.taetigkeiten()}
    for zr, eintraege in sorted(nach_monat.items()):
        if zr in vorhanden:
            ergebnis.uebersprungen.append(zr)
            continue
        for t in eintraege:
            db.taetigkeit_anlegen(t)
        ergebnis.monate.append(zr)
        ergebnis.eintraege += len(eintraege)
    vorlage_speichern(db, v)
    return ergebnis
