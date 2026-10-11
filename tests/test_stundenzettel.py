import io
from datetime import date, datetime
from decimal import Decimal

import openpyxl
import pytest
from fastapi.testclient import TestClient

from ustva import stundenzettel
from ustva.config import Config
from ustva.web import create_app


@pytest.fixture
def client(db, umrechner):
    return TestClient(create_app(Config(bundesland="NW"), db, umrechner, heute=date(2026, 10, 11)))


def tage(*paare):
    return {"tage": [{"datum": d, "stunden": h} for d, h in paare]}


def test_verteilen_in_halben_stunden():
    mittwoche = [date(2026, 9, 2), date(2026, 9, 9), date(2026, 9, 16)]
    assert [h for _, h in stundenzettel.verteilen(Decimal(20), mittwoche)] == [Decimal(7), Decimal("6.5"),
                                                                               Decimal("6.5")]
    assert [h for _, h in stundenzettel.verteilen(Decimal(20), mittwoche[:2])] == [Decimal(10), Decimal(10)]
    assert stundenzettel.verteilen(Decimal(20), []) == []


def test_vorschlag_mittwoche_ohne_feiertage(client):
    v = client.get("/api/stundenzettel/vorschlag/2026-09").json()
    assert v["tage"] == [{"datum": f"2026-09-{t:02d}", "stunden": "4"} for t in (2, 9, 16, 23, 30)]
    assert v["feiertage"] == []
    v = client.get("/api/stundenzettel/vorschlag/2026-02").json()
    assert [t["stunden"] for t in v["tage"]] == ["5"] * 4

    # 2024: Weihnachten fällt auf einen Mittwoch
    v = client.get("/api/stundenzettel/vorschlag/2024-12").json()
    assert [t["datum"] for t in v["tage"]] == ["2024-12-04", "2024-12-11", "2024-12-18"]
    assert [t["stunden"] for t in v["tage"]] == ["7", "6.5", "6.5"]
    assert v["feiertage"] == [{"datum": "2024-12-25", "name": "1. Weihnachtsfeiertag"}]


def test_einstellungen(client):
    assert client.get("/api/stundenzettel/einstellungen").json() == {"name": "", "stunden": "20", "wochentag": 2}
    e = client.put("/api/stundenzettel/einstellungen", json={"name": " Muster Max ", "stunden": "10", "wochentag": 0})
    assert e.json() == {"name": "Muster Max", "stunden": "10", "wochentag": 0}
    v = client.get("/api/stundenzettel/vorschlag/2026-09").json()
    assert [t["datum"] for t in v["tage"]] == ["2026-09-07", "2026-09-14", "2026-09-21", "2026-09-28"]
    assert client.put("/api/stundenzettel/einstellungen", json={"stunden": "0"}).status_code == 400
    assert client.put("/api/stundenzettel/einstellungen", json={"wochentag": 7}).status_code == 400


def test_jahr_bis_zum_letzten_monat(client):
    j = client.get("/api/stundenzettel/2026").json()
    assert j["aktuell"] == "2026-10"
    assert j["vergangen"] == "2026-09"
    assert [m["name"] for m in j["monate"]][-1] == "September"
    assert len(j["monate"]) == 9

    # Der laufende Monat kommt dazu, sobald etwas drinsteht
    r = client.put("/api/stundenzettel/monat/2026-10", json=tage(("2026-10-07", "5")))
    assert r.status_code == 200, r.text
    assert r.json()[-1] == {"zeitraum": "2026-10", "name": "Oktober", "stunden": "5",
                            "tage": [{"datum": "2026-10-07", "stunden": "5"}]}
    assert len(client.get("/api/stundenzettel/2025").json()["monate"]) == 12
    assert client.get("/api/stundenzettel/2027").status_code == 400


def test_monat_eintragen_ersetzt_und_prueft(client):
    client.put("/api/stundenzettel/monat/2026-09", json=tage(("2026-09-02", "5"), ("2026-09-09", "5")))
    # Urlaub: andere Tage, andere Stunden
    client.put("/api/stundenzettel/monat/2026-09", json=tage(("2026-09-01", "10"), ("2026-09-30", "10")))
    sept = client.get("/api/stundenzettel/2026").json()["monate"][8]
    assert sept["tage"] == [{"datum": "2026-09-01", "stunden": "10"}, {"datum": "2026-09-30", "stunden": "10"}]
    assert sept["stunden"] == "20"

    def fehler(zr, daten):
        r = client.put(f"/api/stundenzettel/monat/{zr}", json=daten)
        assert r.status_code == 400
        return r.json()["detail"]

    assert "Zukunft" in fehler("2026-11", tage(("2026-11-04", "5")))
    assert "nicht im August" in fehler("2026-08", tage(("2026-09-02", "5")))
    assert "doppelt" in fehler("2026-09", tage(("2026-09-02", "5"), ("2026-09-02", "5")))
    assert "zwischen 0 und 24" in fehler("2026-09", tage(("2026-09-02", "25")))
    assert client.get("/api/stundenzettel/2026").json()["monate"][8]["stunden"] == "20"

    # Leere Liste leert den Monat
    client.put("/api/stundenzettel/monat/2026-09", json={"tage": []})
    assert client.get("/api/stundenzettel/2026").json()["monate"][8]["tage"] == []


def test_excel_wie_vorlage(client):
    client.put("/api/stundenzettel/einstellungen", json={"name": "Muster Max"})
    client.put("/api/stundenzettel/monat/2026-09", json=tage(*[(f"2026-09-{t:02d}", "4") for t in (2, 9, 16, 23, 30)]))
    acht = tage(*[(f"2026-03-{t:02d}", "2.5") for t in (2, 3, 4, 5, 9, 10, 11, 12)])
    client.put("/api/stundenzettel/monat/2026-03", json=acht)

    r = client.get("/api/stundenzettel/2026/excel")
    assert r.status_code == 200
    assert "Muster%20Max%20Stundenzettel%202026.xlsx" in r.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Januar", "Februar", "März", "April", "Mai", "Juni", "Juli", "August", "September"]

    ws = wb["September"]
    assert ws["A1"].value == "Muster Max"
    assert ws["C1"].value == "Stundenzettel"
    assert ws["C3"].value == "September"
    assert [ws.cell(row=r, column=1).value for r in range(5, 10)] == ["Mittwoch"] * 5
    assert ws["C5"].value == datetime(2026, 9, 2)
    assert ws["C5"].number_format == "dd/mm/yyyy"
    assert [ws.cell(row=r, column=5).value for r in range(5, 10)] == [4] * 5
    assert ws["A13"].value == "Gesamt:"
    assert ws["E13"].value == "=SUM(E5:E11)"
    assert ws["E11"].border.bottom.style == "double"
    assert ws["C17"].value == "Datum" and ws["E17"].value == "Unterschrift"

    # Mehr als sieben Tage: die Summe rutscht nach unten
    ws = wb["März"]
    assert ws["A5"].value == "Montag"
    assert ws["E12"].value == 2.5
    assert ws["A14"].value == "Gesamt:"
    assert ws["E14"].value == "=SUM(E5:E12)"

    leer = wb["Januar"]
    assert leer["A5"].value is None
    assert leer["E13"].value == "=SUM(E5:E11)"


def _alte_datei() -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, datum in (("Januar", [datetime(2025, 1, 8)]), ("März", [datetime(2026, 3, 4), datetime(2026, 3, 11)]),
                        ("April", [datetime(2026, 4, 1)])):
        ws = wb.create_sheet(name)
        ws["A1"], ws["C1"], ws["A3"], ws["C3"], ws["E3"] = "Muster Max", "Stundenzettel", "Monat:", name, "Stunden"
        for i, d in enumerate(datum):
            ws.cell(row=5 + i, column=1, value="Mittwoch")
            ws.cell(row=5 + i, column=3, value=d)
            ws.cell(row=5 + i, column=5, value=5)
        ws["A13"], ws["E13"] = "Gesamt:", "=SUM(E5:E11)"
        ws["C17"], ws["E17"] = "Datum", "Unterschrift"
    puffer = io.BytesIO()
    wb.save(puffer)
    return puffer.getvalue()


def test_import(client):
    client.put("/api/stundenzettel/monat/2026-04", json=tage(("2026-04-08", "20")))
    r = client.post("/api/stundenzettel/import/2026", content=_alte_datei())
    assert r.status_code == 200, r.text
    assert r.json() == {"tage": 2, "monate": ["2026-03"], "uebersprungen": ["2026-04"],
                        "andere_jahre": ["Januar (2025)"], "name": "Muster Max"}
    monate = client.get("/api/stundenzettel/2026").json()["monate"]
    assert monate[2]["stunden"] == "10"
    assert monate[3]["tage"] == [{"datum": "2026-04-08", "stunden": "20"}]
    assert client.get("/api/stundenzettel/einstellungen").json()["name"] == "Muster Max"

    assert client.post("/api/stundenzettel/import/2026", content=b"kein excel").status_code == 400
    leer = io.BytesIO()
    openpyxl.Workbook().save(leer)
    assert client.post("/api/stundenzettel/import/2026", content=leer.getvalue()).status_code == 400
